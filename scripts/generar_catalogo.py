import os
import re
import datetime
import pymongo
import pandas as pd
from bson import ObjectId
from dotenv import load_dotenv

load_dotenv()
MONGO_URI = os.getenv("MONGO_URI") or os.getenv("MONGODB_URI")

BASE_DE_DATOS = "urbanrush_db"
COLECCIONES = [
    "clientes", "productos", "repartidores", "pedidos",
    "pedidos_analiticos",
    "dim_tiempo", "dim_cliente", "dim_producto", "fact_ventas",
]
ROLES = {
    "clientes": "OLTP", "productos": "OLTP", "repartidores": "OLTP", "pedidos": "OLTP",
    "pedidos_analiticos": "OLAP",
    "dim_tiempo": "Data Warehouse", "dim_cliente": "Data Warehouse",
    "dim_producto": "Data Warehouse", "fact_ventas": "Data Warehouse",
}
CAMPOS_FK = {
    "cliente_id", "producto_id", "repartidor_id", "venta_id", "fecha_id",
    "cliente.id", "producto.id", "repartidor.id",
}
MUESTRA = 200

DESCRIPCIONES = {
    "_id": "Identificador unico (ObjectId)",
    "codigo_cliente": "Codigo interno del cliente",
    "nombre": "Nombre del cliente / producto / repartidor segun la coleccion",
    "email": "Correo electronico unico del cliente",
    "telefono": "Numero de telefono de contacto",
    "direccion.calle": "Direccion fisica (calle y numero)",
    "direccion.ciudad": "Ciudad de residencia del cliente",
    "direccion.barrio": "Barrio de residencia del cliente",
    "direccion.pais": "Pais del cliente (Colombia)",
    "fecha_registro": "Fecha en que el cliente se registro",
    "activo": "Indica si el cliente esta activo (booleano)",
    "sku": "Codigo unico de producto (Stock Keeping Unit)",
    "descripcion": "Descripcion comercial del producto",
    "categoria": "Categoria a la que pertenece el producto",
    "precio_unitario": "Precio de venta unitario en COP",
    "stock": "Unidades disponibles en inventario",
    "proveedor": "Proveedor / marca del producto",
    "vendor_id": "Identificador del vendedor asociado al producto",
    "fecha_creacion": "Fecha de creacion del producto",
    "codigo_repartidor": "Codigo interno del repartidor",
    "vehiculo": "Tipo de vehiculo del repartidor (moto, bici, carro)",
    "placa": "Placa del vehiculo",
    "zona": "Zona/barrio de operacion del repartidor",
    "estado": "Estado actual (Entregado, Cancelado, En camino, Pendiente...)",
    "fecha_ingreso": "Fecha de ingreso del repartidor",
    "numero_pedido": "Numero de factura/pedido (negocio)",
    "fecha_pedido": "Fecha y hora del pedido",
    "items": "Detalle de las lineas del pedido (array)",
    "items.producto_id": "FK al producto de la linea",
    "items.nombre_producto": "Nombre del producto al momento de la venta",
    "items.cantidad": "Cantidad vendida de la linea",
    "items.precio_unitario": "Precio unitario de la linea",
    "items.subtotal": "Subtotal de la linea (cantidad * precio)",
    "subtotal": "Suma de los subtotales del pedido",
    "delivery_fee": "Costo de envio en COP",
    "comision_plataforma": "Comision de la plataforma (10% del subtotal)",
    "total": "Total pagado por el cliente (subtotal + delivery)",
    "metodo_pago": "Metodo de pago (Nequi, Tarjeta, PSE, Efectivo...)",
    "ciudad_entrega": "Ciudad donde se entrega el pedido",
    "anio": "Anio de la fecha del pedido (dimension tiempo)",
    "mes": "Mes de la fecha del pedido (1-12)",
    "dia": "Dia del mes de la fecha del pedido",
    "cliente.id": "FK al cliente",
    "cliente.nombre": "Nombre del cliente (desnormalizado)",
    "cliente.ciudad": "Ciudad del cliente (desnormalizado)",
    "cliente.barrio": "Barrio del cliente (desnormalizado)",
    "cliente.segmento": "Segmento del cliente (Premium/Frecuente/Regular/VIP)",
    "producto.id": "FK al producto",
    "producto.nombre": "Nombre del producto (desnormalizado)",
    "producto.categoria": "Categoria del producto (desnormalizado)",
    "repartidor.id": "FK al repartidor",
    "repartidor.nombre": "Nombre del repartidor (desnormalizado)",
    "repartidor.vehiculo": "Vehiculo del repartidor (desnormalizado)",
    "repartidor.zona": "Zona del repartidor (desnormalizado)",
    "cantidad": "Cantidad vendida de la linea",
    "subtotal_linea": "Subtotal de la linea analitica",
    "total_pedido": "Total del pedido al que pertenece la linea",
    "ganancia_estimada": "Ganancia estimada (15% del subtotal de linea)",
    "fecha_id": "Clave de la dimension tiempo (YYYY-MM-DD)",
    "trimestre": "Trimestre del anio (Q1-Q4)",
    "es_fin_semana": "Indica si la fecha es fin de semana (Si/No)",
    "numero_compras": "Cantidad de lineas compradas por el cliente",
    "total_compras": "Gasto acumulado del cliente en COP",
    "ticket_promedio": "Valor promedio por compra",
    "producto_id": "FK al producto (dimension producto)",
    "unidades_vendidas": "Unidades vendidas del producto",
    "ingresos": "Ingresos generados por el producto en COP",
    "numero_pedidos": "Pedidos distintos en los que aparece el producto",
    "venta_id": "Numero de pedido (clave de negocio de la factura)",
    "repartidor_id": "FK al repartidor",
    "cliente_id": "FK al cliente (dimension cliente)",
}


def tipo_estandarizado(valor):
    if isinstance(valor, ObjectId):
        return "ObjectId"
    if isinstance(valor, datetime.datetime):
        return "Timestamp"
    if isinstance(valor, bool):
        return "Boolean"
    if isinstance(valor, int):
        return "Integer"
    if isinstance(valor, float):
        return "Float"
    if isinstance(valor, str):
        return "String"
    if isinstance(valor, list):
        return "Array"
    if isinstance(valor, dict):
        return "Object"
    if valor is None:
        return "Null"
    return type(valor).__name__


def aplanar(prefijo, valor, resultados):
    if isinstance(valor, dict):
        for k, v in valor.items():
            ruta = f"{prefijo}.{k}" if prefijo else k
            aplanar(ruta, v, resultados)
    elif isinstance(valor, list):
        resultados.append((prefijo, "Array"))
        if valor and isinstance(valor[0], dict):
            for k, v in valor[0].items():
                ruta = f"{prefijo}.{k}" if prefijo else k
                aplanar(ruta, v, resultados)
    else:
        resultados.append((prefijo, tipo_estandarizado(valor)))


def clasificar_llave(campo):
    if campo == "_id":
        return "PK"
    if campo in CAMPOS_FK:
        return "FK"
    if re.search(r"_id$", campo) or re.search(r"(^|\.)id$", campo):
        return "FK"
    return ""


def obtener_indices(coleccion):
    indices = set()
    for idx in coleccion.list_indexes():
        for campo in idx["key"]:
            indices.add(campo)
    return indices


def inspeccionar_coleccion(coleccion, nombre):
    indices = obtener_indices(coleccion)
    muestra = list(coleccion.find().limit(MUESTRA))
    total_docs = coleccion.estimated_document_count()

    tipos_por_campo = {}
    presencia = {}
    for doc in muestra:
        campos = []
        aplanar("", doc, campos)
        vistos = set()
        for ruta, tipo in campos:
            vistos.add(ruta)
            tipos_por_campo.setdefault(ruta, {})
            tipos_por_campo[ruta][tipo] = tipos_por_campo[ruta].get(tipo, 0) + 1
            presencia[ruta] = presencia.get(ruta, 0) + 1
    # Campos presentes en todos los docs muestreados o detectados
    for ruta, tipos in tipos_por_campo.items():
        tipo = max(tipos, key=tipos.get)
        pct = (presencia[ruta] / len(muestra)) * 100 if muestra else 0
        yield {
            "coleccion": nombre,
            "campo": ruta,
            "tipo": tipo,
            "llave": clasificar_llave(ruta),
            "indexado": "Si" if ruta in indices else "No",
            "presencia_pct": round(pct, 1),
            "descripcion": DESCRIPCIONES.get(ruta, "Campo de negocio de la coleccion")
        }


def generar_catalogo():
    print("Generando Catalogo de Datos (Data Dictionary)...\n")
    client = pymongo.MongoClient(MONGO_URI, serverSelectionTimeoutMS=20000)
    db = client[BASE_DE_DATOS]

    filas_diccionario = []
    filas_resumen = []

    for nombre in COLECCIONES:
        coleccion = db[nombre]
        total = coleccion.estimated_document_count()
        tamano_mb = db.command("collStats", nombre).get("size", 0) / 1024 / 1024

        campos = list(inspeccionar_coleccion(coleccion, nombre))
        filas_diccionario.extend(campos)
        filas_resumen.append({
            "coleccion": nombre,
            "rol": ROLES[nombre],
            "documentos": total,
            "tamano_MB": round(tamano_mb, 2),
            "numero_campos": len(campos)
        })
        print(f"  [OK] {nombre:<22} {total:>8,} docs | {len(campos):>3} campos | {ROLES[nombre]}")

    df_resumen = pd.DataFrame(filas_resumen)
    df_diccionario = pd.DataFrame(filas_diccionario)

    fecha = datetime.date.today().strftime("%Y%m%d")
    archivo = f"Catalogo_Datos_UrbanRushDB_{fecha}.xlsx"

    with pd.ExcelWriter(archivo, engine="openpyxl") as writer:
        df_resumen.to_excel(writer, sheet_name="Resumen Arquitectonico", index=False)
        df_diccionario.to_excel(writer, sheet_name="Diccionario de Datos", index=False)

    from openpyxl import load_workbook
    from openpyxl.styles import Font

    wb = load_workbook(archivo)
    for ws in wb.worksheets:
        for celda in ws[1]:
            celda.font = Font(bold=True)
        ws.freeze_panes = "A2"
        for col in ws.columns:
            max_len = max((len(str(c.value)) for c in col if c.value is not None), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max(max_len + 2, 12), 60)
    wb.save(archivo)

    print(f"\n[OK] Catalogo generado: {os.path.abspath(archivo)}")
    print(f"  Pestaña 1 'Resumen Arquitectonico': {len(df_resumen)} colecciones")
    print(f"  Pestaña 2 'Diccionario de Datos': {len(df_diccionario)} campos documentados")

    client.close()


if __name__ == "__main__":
    generar_catalogo()